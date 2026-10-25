#!/usr/bin/env python3
"""
Inventar von ZIP-Archiven (nur lesen): Inhalt, Dateitypen, Kopfzeilen und Beispielzeilen.

Gedacht für „Historische Grundfos-Daten – Stammdaten, Bilder, EOL 2015“ (GF_*.zip):
Welche Datei enthält welche Spalten – und welche Sprachen?
Bilder werden nur gezählt. Ergebnis: eine Markdown-Datei im Ausgabeordner.
"""

import collections
import io
import os
import re
import zipfile

from datenaudit import guess_encoding, human, log, looks_like_datanorm

SPRACH_HINWEIS = re.compile(
    r"\b(sprache|language|lang|deu|ger|eng|fra|fre|ita|nld|dut|spa|esp|dan|swe|nor|fin|pol|ces|cze|hun|rus|"
    r"de|en|fr|it|nl|es|da|sv|no|fi|pl|cs|hu|ru)\b|_(de|en|fr|it|nl|es|da|sv|pl|cs|hu|ru)\b", re.I)
BILD = {".jpg", ".jpeg", ".png", ".gif", ".tif", ".tiff", ".bmp", ".eps", ".psd"}


def text_probe(name, data, out, n=3):
    enc = guess_encoding(data[:200000])
    text = data.decode(enc, errors="replace")
    lines = text.splitlines()
    out.append(f"  - `{name}` · {human(len(data))} · {len(lines):,} Zeilen · {enc}"
               + (" · **DATANORM**" if looks_like_datanorm(name, data[:200]) else ""))
    for line in lines[:n]:
        out.append(f"    - `{line[:500]}`")
    if lines:
        treffer = sorted({m.group(0).strip("_").lower() for m in SPRACH_HINWEIS.finditer(lines[0])})
        if treffer:
            out.append(f"    - Sprachhinweise in Kopfzeile: {', '.join(treffer)}")


def xlsx_probe(name, data, out, n=3):
    try:
        import openpyxl
    except ImportError:
        out.append(f"  - `{name}` · {human(len(data))} · (openpyxl fehlt)")
        return
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        out.append(f"  - `{name}` · {human(len(data))} · nicht lesbar: {e!r}")
        return
    out.append(f"  - `{name}` · {human(len(data))} · Blätter: {', '.join(wb.sheetnames)}")
    for ws in wb.worksheets:
        out.append(f"    - Blatt **{ws.title}** ({ws.max_row} × {ws.max_column})")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= n:
                break
            cells = ["" if v is None else str(v) for v in row]
            out.append(f"      - `{' | '.join(cells)[:600]}`")
            if i == 0:
                treffer = sorted({m.group(0).strip("_").lower() for c in cells for m in SPRACH_HINWEIS.finditer(c)})
                if treffer:
                    out.append(f"      - Sprachhinweise: {', '.join(treffer)}")
    wb.close()


def inventar_zip(label, zf, out, tiefe=0):
    infos = [i for i in zf.infolist() if not i.is_dir()]
    exts = collections.Counter(os.path.splitext(i.filename)[1].lower() or "(ohne)" for i in infos)
    out.append(f"- {len(infos)} Dateien, {human(sum(i.file_size for i in infos))} entpackt · "
               + ", ".join(f"{e} {c}" for e, c in exts.most_common(12)))
    for i in infos:
        ext = os.path.splitext(i.filename)[1].lower()
        if ext in BILD:
            continue
        if i.file_size > 400 * 1024 * 1024:
            out.append(f"  - `{i.filename}` · {human(i.file_size)} · zu groß, übersprungen")
            continue
        data = zf.read(i)
        if ext == ".zip" and tiefe < 2:
            out.append(f"  - `{i.filename}` (ZIP im ZIP)")
            try:
                inventar_zip(i.filename, zipfile.ZipFile(io.BytesIO(data)), out, tiefe + 1)
            except zipfile.BadZipFile:
                out.append("    - beschädigt")
        elif ext in (".xlsx", ".xlsm"):
            xlsx_probe(i.filename, data, out)
        elif ext in (".csv", ".txt", ".dat", ".asc", ".xml", ".001", ".002", ".003", ".wrg", ".rab", "") \
                or looks_like_datanorm(i.filename, data[:200]):
            text_probe(i.filename, data, out)
        else:
            out.append(f"  - `{i.filename}` · {human(i.file_size)}")


def run(root, ordner, output_dir="02 Grundfos/00_Datenaudit/ergebnis_gf2015", name="30_zip_inventar.md"):
    base = os.path.join(root, ordner)
    out_dir = os.path.join(root, output_dir)
    os.makedirs(out_dir, exist_ok=True)
    lines = [f"# ZIP-Inventar: {ordner}", ""]
    for fname in sorted(os.listdir(base)):
        path = os.path.join(base, fname)
        if not fname.lower().endswith(".zip"):
            lines.append(f"## {fname} ({human(os.path.getsize(path))}) – keine ZIP, nicht geöffnet\n")
            continue
        log(f"lese {fname} …")
        lines.append(f"## {fname} ({human(os.path.getsize(path))})")
        try:
            with zipfile.ZipFile(path) as zf:
                inventar_zip(fname, zf, lines)
        except zipfile.BadZipFile:
            lines.append("- beschädigt / keine gültige ZIP")
        lines.append("")
    with open(os.path.join(out_dir, name), "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines) + "\n")
    log(f"FERTIG → {os.path.join(out_dir, name)}")
